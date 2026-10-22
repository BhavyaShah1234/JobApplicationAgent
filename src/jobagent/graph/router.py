"""Turn a trigger (the HumanMessage) into work items: job URLs, company names, or rows of a CSV/XLSX file.

- an existing .csv / .xlsx / .xlsm path       → one item per row (company or job URL, optional referral)
- text containing http(s) URLs                → one `url` item per URL
- anything else                               → company names split on commas / semicolons / new lines
                                                ("Acme, Inc." stays one company)

Referrals are attached only with explicit consent (`referral_consent` yes/true/1) and a contact name.
"""

import csv
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Literal
from urllib.parse import urlparse

from jobagent.config import get_settings

URL_RE = re.compile(r"https?://[^\s,;<>\"']+", re.I)
BARE_LINK_RE = re.compile(r"(?<![\w/@.:])(?:www\.)?(?:[a-z0-9-]+\.)+[a-z]{2,}/[^\s,;<>\"']*", re.I)
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")
TABLE_SUFFIXES = {".csv", ".xlsx", ".xlsm"}
# Legal-form tokens that follow a comma inside one company name ("Acme, Inc.").
COMPANY_SUFFIXES = {"inc", "inc.", "llc", "l.l.c.", "ltd", "ltd.", "limited", "corp", "corp.", "corporation",
                    "co", "co.", "company", "plc", "gmbh", "ag", "s.a.", "sa", "b.v.", "bv", "n.v.", "pvt. ltd.",
                    "pvt ltd", "pte. ltd.", "pte ltd", "lp", "l.p.", "llp", "l.l.p.", "pc", "p.c."}
TRUTHY = {"yes", "y", "true", "1", "x", "✓", "consent", "consented"}

# Header aliases (lower-cased, spaces/underscores/hyphens removed).
COLUMNS = {
    "company": {"company", "companyname", "organization", "organisation", "org", "employer", "firm"},
    "url": {"url", "joburl", "link", "joblink", "posting", "postingurl", "applicationurl", "careersurl"},
    "contact_name": {"contactname", "contact", "name", "referrer", "referrername", "referral", "referralname"},
    "contact_email": {"contactemail", "email", "emailaddress", "referreremail", "referralemail"},
    "referral_consent": {"referralconsent", "consent", "referralok", "canrefer", "consenttorefer"},
}


@dataclass
class Referral:
    name: str
    email: str = ""


@dataclass
class WorkItem:
    kind: Literal["url", "company"]
    value: str                     # the URL or the company name
    company: str = ""              # known company (from the file row); '' for bare URLs until scraped
    referral: Referral | None = None
    source: str = ""               # where it came from, e.g. "targets.csv row 3"


@dataclass
class Routed:
    kind: Literal["url", "companies", "file"]
    items: list[WorkItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)


def _norm(header: str) -> str:
    return re.sub(r"[\s_\-]+", "", header.strip().lower())


def _resolve_path(text: str) -> Path | None:
    raw = text.strip().strip("'\"")
    if not raw or "\n" in raw or Path(raw).suffix.lower() not in TABLE_SUFFIXES:
        return None
    s = get_settings()
    for candidate in (Path(raw), s.data_dir / raw, s.data_dir / Path(raw).name):
        if candidate.is_file():
            return candidate
    return None


def split_companies(text: str) -> list[str]:
    parts = [p.strip() for p in re.split(r"[,;\n]+", text) if p.strip()]
    names: list[str] = []
    for part in parts:
        if names and part.lower().strip() in COMPANY_SUFFIXES:
            names[-1] = f"{names[-1]}, {part}"
        else:
            names.append(part)
    seen, unique = set(), []
    for n in names:
        key = re.sub(r"\s+", " ", n).lower()
        if key not in seen:
            seen.add(key)
            unique.append(re.sub(r"\s+", " ", n))
    return unique


def _rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            return [{k or "": (v or "") for k, v in row.items()} for row in csv.DictReader(f, dialect=dialect)]
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = ["" if h is None else str(h) for h in rows[0]]
    return [{header[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(header)} for r in rows[1:]]


def parse_table(path: Path) -> Routed:
    routed = Routed("file")
    rows = _rows(path)
    if not rows:
        routed.warnings.append(f"{path.name}: no rows")
        return routed
    mapping: dict[str, str] = {}
    for header in rows[0]:
        for field_name, aliases in COLUMNS.items():
            if _norm(header) in aliases and field_name not in mapping:
                mapping[field_name] = header
    if "company" not in mapping and "url" not in mapping:
        routed.warnings.append(f"{path.name}: no company or job URL column (headers: {list(rows[0])})")
        return routed
    get = lambda row, f: (row.get(mapping[f], "") if f in mapping else "").strip()  # noqa: E731
    seen: set[tuple[str, str]] = set()
    for n, row in enumerate(rows, start=2):  # row 1 is the header
        where = f"{path.name} row {n}"
        company, url = get(row, "company"), get(row, "url")
        if not company and not url:
            if any(v.strip() for v in row.values()):
                routed.warnings.append(f"{where}: no company or URL, skipped")
            continue
        if url and not URL_RE.fullmatch(url):
            routed.warnings.append(f"{where}: '{url}' isn't a URL, ignored")
            url = ""
        referral = None
        name, email, consent = get(row, "contact_name"), get(row, "contact_email"), get(row, "referral_consent")
        if name or email:
            if consent.lower() not in TRUTHY:
                routed.warnings.append(f"{where}: contact '{name or email}' given without referral consent; "
                                       "not used as a referral")
            elif not name:
                routed.warnings.append(f"{where}: referral consent but no contact name; not used")
            elif email and not EMAIL_RE.match(email):
                routed.warnings.append(f"{where}: invalid email '{email}' for {name}; not used as a referral")
            else:
                referral = Referral(name, email)
        kind, value = ("url", url) if url else ("company", company)
        key = (kind, value.lower())
        if key in seen:
            routed.warnings.append(f"{where}: duplicate of an earlier row, skipped")
            continue
        seen.add(key)
        routed.items.append(WorkItem(kind, value, company=company, referral=referral, source=where))
    return routed


def route(trigger: str) -> Routed:
    """Classify a trigger string and expand it into work items."""
    text = trigger.strip()
    if not text:
        return Routed("companies", warnings=["empty input"])
    path = _resolve_path(text)
    if path is not None:
        return parse_table(path)
    if Path(text.strip("'\"")).suffix.lower() in TABLE_SUFFIXES and "\n" not in text:
        return Routed("file", warnings=[f"file not found: {text} (put it in data/ or give the full path)"])
    # Links pasted without a scheme ("boards.greenhouse.io/acme/jobs/123") count as URLs; bare domains don't.
    text = BARE_LINK_RE.sub(lambda m: "https://" + m.group(0), text)
    urls = [u.rstrip(".)") for u in URL_RE.findall(text)]
    if urls:
        routed = Routed("url")
        seen = set()
        for u in urls:
            if u.lower() not in seen and urlparse(u).netloc:
                seen.add(u.lower())
                routed.items.append(WorkItem("url", u, source="input"))
        leftover = URL_RE.sub(" ", text).strip(" ,;\n")
        if leftover:
            routed.warnings.append(f"ignored non-URL text next to the URLs: {leftover[:80]!r}")
        return routed
    return Routed("companies", [WorkItem("company", c, company=c, source="input") for c in split_companies(text)])


def route_node(state: dict) -> dict:
    """LangGraph node: route the latest HumanMessage."""
    from langchain_core.messages import HumanMessage

    human = next((m for m in reversed(state.get("messages", [])) if isinstance(m, HumanMessage)), None)
    if human is None:
        return {"routed": Routed("companies", warnings=["no HumanMessage to route"]).to_dict()}
    return {"routed": route(str(human.content)).to_dict()}
